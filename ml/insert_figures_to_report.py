"""
Insert real app screenshots and performance charts into the thesis progress report.
Target: BaoCaoTienDo_02-06-2026_NguyenKimNgan.docx
Creates: BaoCaoTienDo_02-06-2026_NguyenKimNgan_with_figures.docx
"""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
REPORTS_DIR  = os.path.join(BASE, "docs", "reports")
SCREENSHOTS  = os.path.join(BASE, "docs", "figures", "screenshots")
REPORT_FIGS  = os.path.join(BASE, "ml", "artifacts", "report_figures")
TSNE_DIR     = os.path.join(BASE, "ml", "artifacts", "tsne")

SRC  = os.path.join(REPORTS_DIR, "BaoCaoTienDo_02-06-2026_NguyenKimNgan.docx")
DEST = os.path.join(REPORTS_DIR, "BaoCaoTienDo_02-06-2026_NguyenKimNgan_with_figures.docx")


def set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def add_heading(doc, text, level=2):
    # Use style_id to avoid duplicate-name KeyError in python-docx
    style_id = f"Heading{level}"
    p = doc.add_paragraph(text, style=style_id)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return p


def add_caption(doc, text, bold_prefix=None):
    """Add an italic caption paragraph below a figure."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if bold_prefix:
        run = p.add_run(bold_prefix)
        run.bold = True
        run.font.size = Pt(9)
    run2 = p.add_run(text)
    run2.italic = True
    run2.font.size = Pt(9)
    return p


def add_figure(doc, img_path, width_inches=5.5, caption=None, caption_prefix=None):
    """Add a centered figure with optional caption."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    try:
        run.add_picture(img_path, width=Inches(width_inches))
    except Exception as e:
        p.add_run(f"[Lỗi chèn ảnh: {e}]")
    if caption:
        add_caption(doc, caption, bold_prefix=caption_prefix)
    return p


def add_two_figures_side_by_side(doc, img1, img2, cap1, cap2, w=3.0):
    """Add two images side by side in a table."""
    tbl = doc.add_table(rows=2, cols=2)
    # No style assignment — use default/no-border table

    # Images
    for i, (img, cap) in enumerate([(img1, cap1), (img2, cap2)]):
        cell = tbl.cell(0, i)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cell.paragraphs[0].add_run()
        try:
            run.add_picture(img, width=Inches(w))
        except Exception as e:
            cell.paragraphs[0].add_run(f"[{e}]")
        # Caption
        cap_cell = tbl.cell(1, i)
        cap_cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap_cell.paragraphs[0].add_run(cap)
        r.italic = True
        r.font.size = Pt(9)
    doc.add_paragraph()


def add_three_figures_row(doc, imgs, captions, w=2.0):
    """Add three images in a row."""
    tbl = doc.add_table(rows=2, cols=3)  # no explicit style
    for i, (img, cap) in enumerate(zip(imgs, captions)):
        cell = tbl.cell(0, i)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cell.paragraphs[0].add_run()
        try:
            run.add_picture(img, width=Inches(w))
        except Exception as e:
            cell.paragraphs[0].add_run(f"[{e}]")
        cap_cell = tbl.cell(1, i)
        cap_cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap_cell.paragraphs[0].add_run(cap)
        r.italic = True
        r.font.size = Pt(9)
    doc.add_paragraph()


def build_figures_section(doc):
    """Append a new section with all figures to the document."""

    # ── Page break before new section
    doc.add_page_break()

    # ── SECTION TITLE
    h = add_heading(doc, "PHỤ LỤC: Hình ảnh minh chứng từ App thực tế", level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph(
        "Tất cả ảnh chụp màn hình được lấy trực tiếp từ emulator Android chạy app "
        "VoiceGuard. Tất cả biểu đồ được vẽ từ số liệu đánh giá thực tế — "
        "không có dữ liệu giả (no fake numbers)."
    )
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    doc.add_paragraph()

    # ══════════════════════════════════════════════════════════════
    # A. APP SCREENSHOTS
    # ══════════════════════════════════════════════════════════════
    add_heading(doc, "A. Ảnh chụp màn hình App (Android Emulator)", level=2)

    # A1. Detect tab — SPOOF result
    doc.add_paragraph().add_run("A.1 — Màn hình Phát hiện: kết quả SPOOF").bold = True
    add_two_figures_side_by_side(
        doc,
        os.path.join(SCREENSHOTS, "02_phat_hien_tab.png"),
        os.path.join(SCREENSHOTS, "04_spoof_details.png"),
        "Kết quả SPOOF (100%) — Phát hiện giả mạo\nBLOCK — Nghi ngờ Deepfake/Replay",
        "Chi tiết: Score 100%, Engine dual-input\nspec+acoustic, Độ trễ 311ms",
        w=2.8
    )

    # A2. Acoustic features + performance
    doc.add_paragraph().add_run("A.2 — Đặc trưng âm thanh và hiệu năng thực tế").bold = True
    add_figure(
        doc,
        os.path.join(SCREENSHOTS, "05_acoustic_features.png"),
        width_inches=2.8,
        caption="8 đặc trưng âm thanh + hiệu năng: 135ms trích đặc trưng, 63ms TFLite inference, 14.7MB RAM",
        caption_prefix="Hình A.2: "
    )

    # A3. Settings benchmark — BONAFIDE + SPOOF side by side
    doc.add_paragraph().add_run("A.3 — Kiểm tra 3 model với file BONAFIDE và SPOOF").bold = True
    add_two_figures_side_by_side(
        doc,
        os.path.join(SCREENSHOTS, "08_settings_benchmark.png"),
        os.path.join(SCREENSHOTS, "09_file2_spoof.png"),
        "File bonafide_VIVOSDEV01_R002.wav\n3 model đều nhận đúng: BONAFIDE (score=0.000234)",
        "File spoof_segment_1109.wav\n3 model đều nhận đúng: SPOOF (score=0.999357)",
        w=2.8
    )

    # A4. Threshold + History side by side
    doc.add_paragraph().add_run("A.4 — Cài đặt ngưỡng và lịch sử phát hiện").bold = True
    add_two_figures_side_by_side(
        doc,
        os.path.join(SCREENSHOTS, "12b_threshold_section.png"),
        os.path.join(SCREENSHOTS, "10_history_tab.png"),
        "Ngưỡng phát hiện = 0.25\nNgưỡng ASV = 0.75",
        "Lịch sử: 4 phiên ghi âm\n4 BLOCK (SPOOF 100%)",
        w=2.8
    )

    doc.add_paragraph()

    # ══════════════════════════════════════════════════════════════
    # B. PERFORMANCE CHARTS
    # ══════════════════════════════════════════════════════════════
    add_heading(doc, "B. Biểu đồ hiệu suất (số liệu từ đánh giá thực tế)", level=2)

    # B1. Confusion matrix
    doc.add_paragraph().add_run("B.1 — Ma trận nhầm lẫn").bold = True
    add_figure(
        doc,
        os.path.join(REPORT_FIGS, "fig_confusion_matrix.png"),
        width_inches=4.5,
        caption="Cross-Scale Attention Lite @ threshold=0.25 | External test: 1.700 mẫu "
                "| TP=830, TN=620, FP=230, FN=20",
        caption_prefix="Hình B.1: "
    )

    # B2. Model comparison
    doc.add_paragraph().add_run("B.2 — So sánh 3 model (Threshold đồng nhất = 0.25)").bold = True
    add_figure(
        doc,
        os.path.join(REPORT_FIGS, "fig_model_comparison.png"),
        width_inches=6.0,
        caption="Cross-Scale đạt F1=86.91%, Recall=97.65% — tốt nhất cho mục tiêu bảo mật "
                "(tối thiểu hóa bỏ sót giả mạo)",
        caption_prefix="Hình B.2: "
    )

    # B3. Threshold calibration
    doc.add_paragraph().add_run("B.3 — Hiệu chỉnh ngưỡng phát hiện").bold = True
    add_figure(
        doc,
        os.path.join(REPORT_FIGS, "fig_threshold_calibration.png"),
        width_inches=6.0,
        caption="Threshold=0.25 được chọn để cân bằng FAR và FRR, tối đa hóa F1 "
                "cho Cross-Scale trên external data",
        caption_prefix="Hình B.3: "
    )

    doc.add_paragraph()

    # ══════════════════════════════════════════════════════════════
    # C. ABLATION STUDY CHARTS
    # ══════════════════════════════════════════════════════════════
    add_heading(doc, "C. Ablation Study — Phân tích 8 đặc trưng âm thanh", level=2)

    # C1. Ablation bar chart
    doc.add_paragraph().add_run("C.1 — Mức độ quan trọng của từng đặc trưng (Leave-one-out)").bold = True
    add_figure(
        doc,
        os.path.join(REPORT_FIGS, "fig_ablation_study.png"),
        width_inches=6.5,
        caption="ActiveDuration quan trọng nhất (F1 giảm 0.48% khi xóa). "
                "DynamicRange ít quan trọng nhất (F1 tăng 0.02% — có thể dư thừa nhỏ).",
        caption_prefix="Hình C.1: "
    )

    # C2. Progressive addition
    doc.add_paragraph().add_run("C.2 — Hiệu năng khi thêm dần đặc trưng (Progressive Addition)").bold = True
    add_figure(
        doc,
        os.path.join(REPORT_FIGS, "fig_progressive_addition.png"),
        width_inches=6.0,
        caption="Accuracy đạt ~98% từ 5 đặc trưng trở lên. Thêm 3 đặc trưng còn lại "
                "cải thiện thêm ~0.02% — giữ đủ 8 để đảm bảo tính tổng quát.",
        caption_prefix="Hình C.2: "
    )

    doc.add_paragraph()

    # ══════════════════════════════════════════════════════════════
    # D. T-SNE VISUALIZATION
    # ══════════════════════════════════════════════════════════════
    add_heading(doc, "D. Trực quan hóa t-SNE — Phân tích domain shift", level=2)

    p = doc.add_paragraph(
        "t-SNE (t-distributed Stochastic Neighbor Embedding) chiếu không gian 8 chiều "
        "của đặc trưng âm thanh xuống 2D để trực quan hóa sự phân tách giữa "
        "bonafide/spoof và sự khác biệt giữa dữ liệu nội bộ và ngoài (domain shift)."
    )
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # D1. Combined t-SNE
    doc.add_paragraph().add_run("D.1 — Tổng hợp 4 góc nhìn t-SNE").bold = True
    tsne_combined = os.path.join(TSNE_DIR, "tsne_combined.png")
    if os.path.exists(tsne_combined):
        add_figure(
            doc,
            tsne_combined,
            width_inches=6.5,
            caption="(Trái) Phân tách bonafide/spoof rõ ràng trong không gian đặc trưng. "
                    "(Phải) Domain shift: dữ liệu internal và external tạo cụm riêng biệt "
                    "→ lý do accuracy giảm 50% trước khi mix external vào training.",
            caption_prefix="Hình D.1: "
        )

    # D2. Label and Source side by side
    tsne_label  = os.path.join(TSNE_DIR, "tsne_label.png")
    tsne_source = os.path.join(TSNE_DIR, "tsne_source.png")
    if os.path.exists(tsne_label) and os.path.exists(tsne_source):
        doc.add_paragraph().add_run("D.2 — Chi tiết: nhãn bonafide/spoof và nguồn dữ liệu").bold = True
        add_two_figures_side_by_side(
            doc,
            tsne_label,
            tsne_source,
            "Màu theo nhãn: bonafide (xanh) vs spoof (đỏ)\n→ Hai nhóm tách biệt rõ",
            "Màu theo nguồn: internal (xanh) vs external (cam)\n→ Domain shift rõ ràng",
            w=3.0
        )

    doc.add_paragraph()

    # ══════════════════════════════════════════════════════════════
    # E. DEVICE PERFORMANCE
    # ══════════════════════════════════════════════════════════════
    add_heading(doc, "E. Hiệu năng trên thiết bị di động", level=2)

    add_figure(
        doc,
        os.path.join(REPORT_FIGS, "fig_device_performance.png"),
        width_inches=6.5,
        caption="Đo từ app chạy trên Android Emulator: tổng pipeline 311ms, "
                "TFLite inference chỉ 63ms (20%), RAM 14.7MB. "
                "Phù hợp real-time trên thiết bị di động.",
        caption_prefix="Hình E.1: "
    )

    print("[OK] Figures section built successfully.")


def main():
    print(f"Loading: {SRC}")
    doc = Document(SRC)

    build_figures_section(doc)

    doc.save(DEST)
    size_kb = os.path.getsize(DEST) // 1024
    print(f"[OK] Saved: {DEST} ({size_kb} KB)")


if __name__ == "__main__":
    main()
